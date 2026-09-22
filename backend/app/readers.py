from __future__ import annotations
import csv
from dataclasses import dataclass
from pathlib import Path
from typing import Iterator
from xml.etree.ElementTree import iterparse
from openpyxl import load_workbook
import xlrd

@dataclass(frozen=True)
class FileProfile:
    columns: list[str]
    preview: list[dict[str, str | None]]
    rows_estimated: int | None

def _clean(value: object) -> str | None:
    if value is None: return None
    return str(value).strip()

class FileReader:
    extensions: tuple[str, ...] = ()
    def profile(self, path: Path, options: dict) -> FileProfile: raise NotImplementedError
    def rows(self, path: Path, options: dict) -> Iterator[dict[str, str | None]]: raise NotImplementedError

class DelimitedReader(FileReader):
    extensions = (".csv", ".txt")
    def _dialect(self, path: Path, options: dict):
        encoding = options.get("encoding", "utf-8-sig")
        delimiter = options.get("delimiter")
        with path.open("r", encoding=encoding, newline="") as handle:
            sample = handle.read(8192)
        if not delimiter:
            delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
        return encoding, delimiter
    def rows(self, path: Path, options: dict):
        encoding, delimiter = self._dialect(path, options)
        with path.open("r", encoding=encoding, newline="") as handle:
            reader = csv.reader(handle, delimiter=delimiter)
            headers = None
            header_row = int(options.get("header_row", 1))
            for index, values in enumerate(reader, 1):
                if index < header_row: continue
                if headers is None:
                    headers = [str(v).strip() or f"column_{i}" for i, v in enumerate(values, 1)]
                    continue
                yield {headers[i]: _clean(values[i]) if i < len(values) else None for i in range(len(headers))}
    def profile(self, path, options):
        sample = []
        iterator = self.rows(path, options)
        for _ in range(int(options.get("preview_rows", 100))):
            try: sample.append(next(iterator))
            except StopIteration: break
        columns = list(sample[0].keys()) if sample else []
        return FileProfile(columns, sample, sum(1 for _ in self.rows(path, options)))

class ExcelReader(FileReader):
    extensions = (".xlsx", ".xls")
    def rows(self, path: Path, options: dict):
        if path.suffix.lower() == ".xls":
            workbook = xlrd.open_workbook(path.as_posix(), on_demand=True)
            sheet = options.get("sheet") or workbook.sheet_names()[0]
            ws = workbook.sheet_by_name(sheet)
            header_row = int(options.get("header_row", 1))
            headers = None
            try:
                for index in range(header_row - 1, ws.nrows):
                    values = ws.row_values(index)
                    if headers is None:
                        headers = [str(v).strip() or f"column_{i}" for i, v in enumerate(values, 1)]
                        continue
                    yield {headers[i]: _clean(values[i]) if i < len(values) else None for i in range(len(headers))}
            finally:
                workbook.release_resources()
            return
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = options.get("sheet") or workbook.sheetnames[0]
        ws = workbook[sheet]; header_row = int(options.get("header_row", 1)); headers = None
        try:
            for index, values in enumerate(ws.iter_rows(values_only=True), 1):
                if index < header_row: continue
                values = list(values)
                if headers is None:
                    headers = [str(v).strip() if v is not None else f"column_{i}" for i, v in enumerate(values, 1)]
                    continue
                yield {headers[i]: _clean(values[i]) if i < len(values) else None for i in range(len(headers))}
        finally:
            workbook.close()
    def profile(self, path, options):
        iterator = self.rows(path, options); sample = []
        for _ in range(int(options.get("preview_rows", 100))):
            try: sample.append(next(iterator))
            except StopIteration: break
        return FileProfile(list(sample[0].keys()) if sample else [], sample, None)

class XmlReader(FileReader):
    extensions = (".xml",)
    def rows(self, path: Path, options: dict):
        record_tag = options.get("record_tag")
        if not record_tag: raise ValueError("XML requiere record_tag para evitar cargar una estructura ambigua")
        for _, element in iterparse(path, events=("end",)):
            if element.tag.rsplit("}", 1)[-1] != record_tag: continue
            record = {child.tag.rsplit("}", 1)[-1]: _clean(child.text) for child in list(element)}
            yield record; element.clear()
    def profile(self, path, options):
        iterator = self.rows(path, options); sample = []
        for _ in range(int(options.get("preview_rows", 100))):
            try: sample.append(next(iterator))
            except StopIteration: break
        return FileProfile(list(sample[0].keys()) if sample else [], sample, None)

READERS = {ext: reader for reader in (DelimitedReader(), ExcelReader(), XmlReader()) for ext in reader.extensions}

def get_reader(filename: str) -> FileReader:
    suffix = Path(filename).suffix.lower()
    if suffix not in READERS: raise ValueError(f"Formato no soportado: {suffix or 'sin extensión'}")
    return READERS[suffix]
